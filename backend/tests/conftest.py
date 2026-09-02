import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.core.auth import require_current_user
from app.core.config import Settings
from app.main import app, create_app
from app.schemas.auth import SessionUser


@pytest.fixture
def authenticated_client() -> TestClient:
    app.dependency_overrides[require_current_user] = lambda: SessionUser(
        user_id="test-user",
        user_name="Test User",
        user_type="研究生",
    )
    try:
        with TestClient(app) as test_client:
            yield test_client
    finally:
        app.dependency_overrides.pop(require_current_user, None)


@pytest.fixture
def anonymous_client() -> TestClient:
    application = create_app(
        Settings(
            _env_file=None,
            app_env="test",
            uis_oauth_client_id="client-id",
            uis_oauth_client_secret="client-secret",
            auth_session_secret="s" * 32,
        )
    )
    with TestClient(application) as test_client:
        yield test_client


@pytest.fixture
def sample_docx() -> bytes:
    document = Document()
    document.add_heading("学工例会备忘", level=1)
    document.add_paragraph("一、学生工作")
    document.add_paragraph("1、教育管理工作")
    document.add_paragraph("（1）【活动报名】请同学们于8月20日（周四）17:00前报名。联系人：张老师 123456。")
    document.add_paragraph("（2）【材料提交】毕业生请于8月21日提交申请材料。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "材料"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "申请表"
    table.cell(1, 1).text = "1份"
    document.add_paragraph("二、工作提醒")
    document.add_paragraph("本段不应进入表格解析结果。")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture
def multi_column_table_docx() -> bytes:
    document = Document()
    document.add_paragraph("一、学生工作")
    document.add_paragraph("1、材料汇总")
    document.add_paragraph("（1）请核对下列表格。")
    table = document.add_table(rows=4, cols=4)
    for row_index, row in enumerate(table.rows, start=1):
        for column_index, cell in enumerate(row.cells, start=1):
            cell.text = f"R{row_index}C{column_index}"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture
def merged_table_docx() -> bytes:
    document = Document()
    document.add_paragraph("一、学生工作")
    document.add_paragraph("1、材料汇总")
    document.add_paragraph("（1）请核对合并单元格。")
    table = document.add_table(rows=4, cols=4)
    for row_index, row in enumerate(table.rows, start=1):
        for column_index, cell in enumerate(row.cells, start=1):
            cell.text = f"R{row_index}C{column_index}"
    table.cell(1, 3).merge(table.cell(3, 3))
    table.cell(3, 0).merge(table.cell(3, 1))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
