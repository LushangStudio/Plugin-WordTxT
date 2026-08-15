from __future__ import annotations

import html
import io
import re
import uuid
from dataclasses import dataclass, field

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.schemas.models import (
    CategoryMarker,
    CutInfo,
    ExtractDocumentResponse,
    ExtractedDocumentTable,
    ParseDocumentResponse,
    ParsedItem,
    TableReport,
)


CATEGORY_RE = re.compile(
    r"^\s*([一二三四五六七八九十百千]+)、\s*([^\n【】#:：]{2,50})\s*$", re.MULTILINE
)
SPECIAL_CATEGORY_RE = re.compile(r"^\s*(专项工作)\s*$", re.MULTILINE)
SUBCATEGORY_RE = re.compile(
    r"^\s*((?:[1-9]|1[0-9]))[、\.．\)]\s*([^\n【】#:：]{1,100})\s*$", re.MULTILINE
)
ITEM_MARKER_RE = re.compile(r"（\s*\d+\s*）|\(\s*\d+\s*\)")
DATE_RE = re.compile(
    r"(?P<date>\d{1,2}月\d{1,2}日)(?:（(?P<week>周[一二三四五六日天]|星期[一二三四五六日天])）)?"
)
TIME_RE = re.compile(r"((?:上|下)午)?\s*(\d{1,2})[：:](\d{2})")
CUT_HEADING_RE = re.compile(
    r"^\s*(?:(?:[一二三四五六七八九十百千]+|\d+)[、\.．\)）]\s*)?(工作提醒|发放材料)(?:[:：])?\s*$",
    re.MULTILINE,
)


@dataclass
class ExtractedTable:
    token: str
    html: str
    plain: str
    position: int


@dataclass
class LinearDocument:
    text: str
    tables: list[ExtractedTable] = field(default_factory=list)
    report: TableReport = field(default_factory=TableReport)


@dataclass
class StructureItem:
    text: str
    position: int
    end: int = 0


def normalize_memo_text(value: str) -> str:
    text = str(value or "")
    text = text.replace("\u00a0", " ").replace("\u3000", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\[[^\]\n]{1,20}\d+(?:\.\d+)+\]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _iter_block_items(document: DocumentObject):
    """按 XML 文档顺序递归读取段落和表格。

    学工备忘常由 WPS 生成，正文可能包在 ``w:sdt`` 等内容控件中，
    只遍历 body 的直接子节点会漏掉绝大部分文本。
    """

    def visit(node):
        for child in node.iterchildren():
            if child.tag == qn("w:tbl"):
                yield Table(child, document)
                continue
            if child.tag == qn("w:p"):
                yield Paragraph(child, document)
                continue
            yield from visit(child)

    yield from visit(document.element.body)


def _paragraph_html(paragraph: Paragraph) -> str:
    pieces: list[str] = []
    for run in paragraph.runs:
        value = html.escape(run.text or "")
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        pieces.append(value)
    alignment = str(paragraph.alignment or "")
    css = "text-align:center;" if "CENTER" in alignment else "text-align:right;" if "RIGHT" in alignment else ""
    return f'<div style="{css}">{"".join(pieces)}</div>'


def _table_to_html(table: Table) -> str:
    cell_rows: list[list[tuple[_Cell, int, int, str | None]]] = []
    for row in table.rows:
        grid_column = row.grid_cols_before
        row_cells: list[tuple[_Cell, int, int, str | None]] = []
        # Use the direct XML cells rather than ``row.cells``.  The latter
        # expands horizontally/vertically merged cells into repeated wrappers,
        # which makes it impossible to preserve the original spans reliably.
        for cell_element in row._tr.tc_lst:
            column_span = cell_element.grid_span
            row_cells.append(
                (
                    _Cell(cell_element, table),
                    grid_column,
                    column_span,
                    cell_element.vMerge,
                )
            )
            grid_column += column_span
        cell_rows.append(row_cells)

    rows: list[str] = []
    for row_index, row_cells in enumerate(cell_rows):
        cells: list[str] = []
        for cell, grid_column, column_span, vertical_merge in row_cells:
            if vertical_merge == "continue":
                continue

            row_span = 1
            if vertical_merge == "restart":
                for following_row in cell_rows[row_index + 1 :]:
                    continuation = next(
                        (
                            candidate
                            for candidate in following_row
                            if candidate[1] == grid_column
                        ),
                        None,
                    )
                    if continuation is None or continuation[3] != "continue":
                        break
                    row_span += 1

            content = "".join(_paragraph_html(paragraph) for paragraph in cell.paragraphs)
            nested = "".join(_table_to_html(nested_table) for nested_table in cell.tables)
            span_attributes = ""
            if column_span > 1:
                span_attributes += f' colspan="{column_span}"'
            if row_span > 1:
                span_attributes += f' rowspan="{row_span}"'
            cells.append(
                f'<td{span_attributes} style="border:1px solid #d1d5db;padding:8px;vertical-align:top;">'
                + (content + nested or "&nbsp;")
                + "</td>"
            )
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return (
        '<table class="docx-table" style="border-collapse:collapse;width:100%;margin:12px 0;">'
        "<tbody>" + "".join(rows) + "</tbody></table>"
    )


def _table_plain_text(table: Table) -> str:
    lines: list[str] = []
    for row in table.rows:
        values = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
        if any(values):
            lines.append(" | ".join(values))
    return "\n".join(lines)


def read_docx(payload: bytes) -> LinearDocument:
    document = Document(io.BytesIO(payload))
    parts: list[str] = []
    tables: list[ExtractedTable] = []
    report = TableReport()
    table_index = 0

    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                parts.append(text)
            continue

        table_index += 1
        report.found += 1
        try:
            token = f"<!--DOCX_TABLE_{table_index}-->"
            position = len("\n".join(parts)) + (1 if parts else 0)
            tables.append(
                ExtractedTable(
                    token=token,
                    html=_table_to_html(block),
                    plain=_table_plain_text(block),
                    position=position,
                )
            )
            parts.append(token)
            report.converted += 1
        except Exception:
            report.failed += 1

    return LinearDocument(text=normalize_memo_text("\n".join(parts)), tables=tables, report=report)


def _replace_table_tokens_with_plain(text: str, tables: list[ExtractedTable]) -> str:
    output = text
    for table in tables:
        output = output.replace(table.token, table.plain)
    return normalize_memo_text(output)


def extract_document(filename: str, payload: bytes) -> ExtractDocumentResponse:
    linear = read_docx(payload)
    plain_text = _replace_table_tokens_with_plain(linear.text, linear.tables)
    text = linear.text
    extracted_tables: list[ExtractedDocumentTable] = []
    for index, table in enumerate(linear.tables, start=1):
        marker = f"[[表格_{index}]]"
        text = text.replace(table.token, marker)
        extracted_tables.append(
            ExtractedDocumentTable(
                marker=marker,
                html=table.html,
                plain_text=table.plain,
            )
        )
    return ExtractDocumentResponse(
        filename=filename,
        text=normalize_memo_text(text),
        character_count=len(plain_text),
        tables=extracted_tables,
        table_report=linear.report,
    )


def _find_cut(text: str) -> tuple[str, CutInfo]:
    match = CUT_HEADING_RE.search(text)
    if not match:
        return text, CutInfo()
    paragraph_index = text[: match.start()].count("\n") + 1
    return text[: match.start()].rstrip(), CutInfo(
        keyword=match.group(1),
        paragraph_index=paragraph_index,
        offset=match.start(),
    )


def _find_markers(text: str) -> list[CategoryMarker]:
    markers: list[CategoryMarker] = []
    for match in CATEGORY_RE.finditer(text):
        markers.append(
            CategoryMarker(
                title=f"{match.group(1)}、{match.group(2).strip()}",
                position=match.start(),
                level=1,
            )
        )
    for match in SPECIAL_CATEGORY_RE.finditer(text):
        markers.append(CategoryMarker(title=match.group(1), position=match.start(), level=1))
    for match in SUBCATEGORY_RE.finditer(text):
        markers.append(
            CategoryMarker(
                title=f"{match.group(1)}、{match.group(2).strip()}",
                position=match.start(),
                level=2,
            )
        )
    markers.sort(key=lambda marker: marker.position)
    return markers


def _marker_line_end(text: str, position: int) -> int:
    newline = text.find("\n", position)
    return len(text) if newline < 0 else newline + 1


def _should_keep_section_single(raw: str) -> bool:
    match = ITEM_MARKER_RE.search(raw)
    if not match or match.start() <= 0:
        return False
    before = raw[: match.start()]
    return bool(re.search(r"[。；;：:]\s*$", before) or len(before) > 20)


def _split_section_items(raw: str, base_position: int) -> list[StructureItem]:
    stripped = raw.strip()
    if not stripped:
        return []
    leading = raw.find(stripped)
    matches = list(ITEM_MARKER_RE.finditer(stripped))
    if not matches or matches[0].start() != 0 or _should_keep_section_single(stripped):
        return [StructureItem(text=stripped, position=base_position + leading)]

    items: list[StructureItem] = []
    for index, match in enumerate(matches):
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(stripped)
        chunk = stripped[start:end].strip()
        if chunk:
            items.append(StructureItem(text=chunk, position=base_position + leading + start))
    return items


def _parse_structure(text: str) -> tuple[list[CategoryMarker], list[StructureItem]]:
    markers = _find_markers(text)
    items: list[StructureItem] = []
    for index, marker in enumerate(markers):
        following = markers[index + 1 :]
        next_marker = next(
            (
                candidate
                for candidate in following
                if candidate.level <= marker.level or marker.level == 2
            ),
            None,
        )
        start = _marker_line_end(text, marker.position)
        end = next_marker.position if next_marker else len(text)
        if marker.level == 1:
            first_sub = next(
                (
                    candidate
                    for candidate in following
                    if candidate.level == 2 and candidate.position < end
                ),
                None,
            )
            if first_sub:
                end = first_sub.position
        items.extend(_split_section_items(text[start:end], start))

    if not items:
        block_matches = list(re.finditer(r"^\s*(?:\d+[、\.．]?\s*)?【[^】]{1,80}】", text, re.MULTILINE))
        for index, match in enumerate(block_matches):
            end = block_matches[index + 1].start() if index + 1 < len(block_matches) else len(text)
            chunk = text[match.start() : end].strip()
            if chunk:
                items.append(StructureItem(text=chunk, position=match.start()))

    items.sort(key=lambda item: item.position)
    for index, item in enumerate(items):
        item.end = items[index + 1].position if index + 1 < len(items) else len(text)
    return markers, items


def _extract_time(text: str) -> tuple[str, str]:
    date_match = DATE_RE.search(text)
    date = date_match.group("date") if date_match else ""
    weekday = date_match.group("week") if date_match and date_match.group("week") else ""
    times: list[str] = []
    for match in TIME_RE.finditer(text):
        value = f"{match.group(1) or ''}{int(match.group(2)):02d}:{match.group(3)}"
        if value not in times:
            times.append(value)
    if date and times:
        return f"{date} {'、'.join(times)}", weekday
    return date or "、".join(times), weekday


def _extract_contact(text: str) -> tuple[str, str]:
    match = re.search(r"联系人[:：]\s*([\s\S]*?)(?:。|\n|$)", text)
    if not match:
        return "", ""
    names: list[str] = []
    phones: list[str] = []
    for part in re.split(r"[，,]", match.group(1)):
        pair = re.match(r"^([^\s，。]+)\s+([0-9\-－—]+)$", part.strip())
        if pair:
            names.append(pair.group(1))
            phones.append(pair.group(2))
    return "、".join(names), "，".join(phones)


def _summary(text: str) -> str:
    clean = re.sub(r"<!--DOCX_TABLE_\d+-->", "", text)
    clean = re.sub(r"^\s*(?:（\s*\d+\s*）|\(\s*\d+\s*\)|\d+[、\.．\)）])\s*", "", clean)
    clean = re.sub(r"\s+", " ", clean).strip()
    sentence = re.match(r"^[\s\S]*?[。！？!?](?=\s|$|[\u4e00-\u9fff])", clean)
    if sentence:
        return sentence.group(0).strip()
    return clean if len(clean) <= 80 else clean[:80] + "..."


def _content_html(raw: str, tables: list[ExtractedTable]) -> str:
    output = html.escape(raw).replace("\n", "<br>")
    for table in tables:
        output = output.replace(html.escape(table.token), table.html)
    return output


def _category_context(markers: list[CategoryMarker], position: int) -> tuple[str, str]:
    category = ""
    subcategory = ""
    for marker in markers:
        if marker.position > position:
            break
        if marker.level == 1:
            category = marker.title
            subcategory = ""
        else:
            subcategory = marker.title
    return category, subcategory


def parse_document(filename: str, payload: bytes, start_number: int = 1) -> ParseDocumentResponse:
    linear = read_docx(payload)
    kept_text, cut_info = _find_cut(linear.text)
    kept_tables = [table for table in linear.tables if table.position < len(kept_text)]
    markers, structure_items = _parse_structure(kept_text)
    parsed_items: list[ParsedItem] = []

    for index, item in enumerate(structure_items):
        item_tables = [
            table
            for table in kept_tables
            if table.token in item.text or item.position <= table.position < item.end
        ]
        clean = item.text
        for table in item_tables:
            clean = clean.replace(table.token, "")
        clean = normalize_memo_text(clean)
        time_value, weekday = _extract_time(clean)
        contact_name, contact_phone = _extract_contact(clean)
        category, subcategory = _category_context(markers, item.position)
        parsed_items.append(
            ParsedItem(
                uid=str(uuid.uuid4()),
                number=start_number + index,
                time=time_value,
                weekday=weekday,
                contact_name=contact_name,
                contact_phone=contact_phone,
                summary=_summary(clean),
                content=clean,
                content_html=_content_html(item.text, item_tables),
                tables_html=[table.html for table in item_tables],
                raw_position=item.position,
                category=category,
                subcategory=subcategory,
            )
        )

    return ParseDocumentResponse(
        filename=filename,
        items=parsed_items,
        categories=markers,
        cut_info=cut_info,
        table_report=TableReport(
            found=linear.report.found,
            converted=len(kept_tables),
            failed=linear.report.failed,
        ),
    )
