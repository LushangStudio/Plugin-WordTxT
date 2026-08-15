from __future__ import annotations

import io
import math
import unicodedata
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.models import ExportRequest


HEADERS = ["编号", "时间", "星期几", "联系人", "联系电话", "工作概述", "提取的内容", "备注"]
COLUMN_WIDTH_LIMITS = [
    (8, 10),
    (14, 24),
    (10, 12),
    (12, 20),
    (16, 24),
    (20, 42),
    (36, 72),
    (14, 32),
]
METADATA_WIDTH_LIMITS = [(12, 24), (24, 60)]
HEADER_HEIGHT = 30
MIN_BODY_ROW_HEIGHT = 24
MAX_BODY_ROW_HEIGHT = 409
BODY_LINE_HEIGHT = 15
HEADER_FILL = PatternFill(fill_type="solid", fgColor="2457D6")
BODY_DIVIDER = Border(bottom=Side(style="thin", color="E5EAF1"))


def _character_display_width(character: str) -> int:
    if character == "\t":
        return 4
    if unicodedata.combining(character):
        return 0
    return 2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1


def _line_display_width(value: object) -> int:
    lines = str(value or "").splitlines() or [""]
    return max(sum(_character_display_width(character) for character in line) for line in lines)


def _wrapped_line_count(value: object, column_width: float) -> int:
    lines = str(value or "").splitlines() or [""]
    # Excel column-width units leave slightly more room than a strict
    # character count, especially with proportional Latin fonts.  The factor
    # keeps estimated row heights close to Excel's native AutoFit result while
    # East Asian full-width characters are still counted as two units.
    usable_width = max(1, int(column_width * 1.1) - 2)
    return sum(
        max(
            1,
            math.ceil(
                sum(_character_display_width(character) for character in line)
                / usable_width
            ),
        )
        for line in lines
    )


def _fit_worksheet(
    sheet: Worksheet,
    width_limits: list[tuple[float, float]],
) -> None:
    column_widths: dict[int, float] = {}
    for column_index, (minimum, maximum) in enumerate(width_limits, start=1):
        measured_width = max(
            (_line_display_width(sheet.cell(row=row, column=column_index).value) for row in range(1, sheet.max_row + 1)),
            default=0,
        )
        width = min(maximum, max(minimum, measured_width + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        column_widths[column_index] = width

    sheet.row_dimensions[1].height = HEADER_HEIGHT
    for row_index in range(2, sheet.max_row + 1):
        required_lines = max(
            _wrapped_line_count(
                sheet.cell(row=row_index, column=column_index).value,
                column_widths[column_index],
            )
            for column_index in column_widths
        )
        sheet.row_dimensions[row_index].height = min(
            MAX_BODY_ROW_HEIGHT,
            max(MIN_BODY_ROW_HEIGHT, required_lines * BODY_LINE_HEIGHT + 6),
        )


def build_workbook(request: ExportRequest) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "通知整理"
    sheet.append(HEADERS)

    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", size=12, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)

    for item in request.items:
        sheet.append(
            [
                item.number,
                item.time,
                item.weekday,
                item.contact_name,
                item.contact_phone,
                item.summary,
                item.content,
                item.notes,
            ]
        )

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=11)
            cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
            cell.border = BODY_DIVIDER

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    _fit_worksheet(sheet, COLUMN_WIDTH_LIMITS)

    metadata = workbook.create_sheet("元数据")
    metadata.append(["键", "值"])
    metadata.append(["排序方式", request.sort_type])
    metadata.append(["文档名", request.source_document])
    metadata.append(["导出时间", datetime.now().isoformat(timespec="seconds")])
    metadata.append(["记录条数", len(request.items)])
    for cell in metadata[1]:
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="left")
    for row in metadata.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
            cell.border = BODY_DIVIDER
    metadata.freeze_panes = "A2"
    metadata.sheet_view.showGridLines = False
    _fit_worksheet(metadata, METADATA_WIDTH_LIMITS)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
